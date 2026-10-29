"""Merge research results and build the material/equipment passport table.

Usage: python scripts/build_table.py <research_dir>
Outputs: data/catalog.json,
         output/materials_equipment.xlsx (v1: normalized PDF file names),
         output/materials_equipment_links.xlsx (v2: direct PDF hyperlinks)
Main sheets list only PDFs that were downloaded and content-verified; the
"Manbalar" sheet keeps every source with its status.
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlparse, urlsplit, urlunsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
MATERIALS_FILE = ROOT / "data" / "materials.json"
CATALOG_FILE = ROOT / "data" / "catalog.json"
SELECTION_FILE = ROOT / "data" / "equipment_selection.json"
XLSX_FILE = ROOT / "output" / "materials_equipment.xlsx"
XLSX_LINKS_FILE = ROOT / "output" / "materials_equipment_links.xlsx"
DOWNLOAD_REPORT = ROOT / "data" / "download_report.json"

MATERIAL_DIR = "pdf/materials"
EQUIPMENT_DIR = "pdf/equipment"
PDF_URL_HINTS = ("getdocument", "/purl/", "msds.ford.com/documents", "pdf_main.aspx")
NO_PASSPORT_KEYS = {"WIPE"}
LINK_FONT = Font(color="0563C1", underline="single")


def slug(text: str, max_len: int = 60) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", ascii_text).strip("_")
    return cleaned[:max_len].rstrip("_") or "NA"


def is_pdf_url(url: str) -> bool:
    lowered = url.lower()
    return urlparse(lowered).path.endswith(".pdf") or any(h in lowered for h in PDF_URL_HINTS)


class FileNamer:
    """Assign one normalized file name per unique URL; HTML pages get no file."""

    def __init__(self) -> None:
        self._by_url: dict[str, str] = {}
        self._taken: set[str] = set()

    def assign(self, url: str, base: str, folder: str) -> str | None:
        if not is_pdf_url(url):
            return None
        if url in self._by_url:
            return self._by_url[url]
        name, seq = f"{base}.pdf", 1
        while name in self._taken:
            seq += 1
            name = f"{base}_{seq}.pdf"
        self._taken.add(name)
        self._by_url[url] = f"{folder}/{name}"
        return self._by_url[url]


def material_base(item: dict[str, Any], doc: dict[str, Any]) -> str:
    """M<id>_<TYPE>[_EQUIV]_<real document title>: groups by material, names the document."""
    exactness = "" if doc.get("exact_or_equivalent", "exact") == "exact" else "_EQUIV"
    title = doc.get("title") or item["name"]
    return f"M{item['id']:02d}_{slug(doc.get('type', 'DOC'), 10)}{exactness}_{slug(title, 70)}"


def equipment_base(doc: dict[str, Any]) -> str:
    title = doc.get("title") or f"{doc.get('vendor', '')} {doc.get('model', '')}"
    return f"EQ_{slug(doc.get('type', 'DOC'), 12)}_{slug(title, 70)}"


def load_research(research_dir: Path) -> dict[int, dict[str, Any]]:
    merged: dict[int, dict[str, Any]] = {}
    for path in sorted(research_dir.glob("group_*.json")):
        for entry in json.loads(path.read_text(encoding="utf-8")):
            merged[int(entry["id"])] = entry
    return merged


def load_selection() -> dict[str, Any]:
    selection = json.loads(SELECTION_FILE.read_text(encoding="utf-8"))
    unknown = {key for keys in selection["materials"].values() for key in keys} - set(selection["equipment"])
    if unknown:
        raise SystemExit(f"equipment_selection.json: unknown equipment keys {sorted(unknown)}")
    return selection


def build_catalog(research_dir: Path, selection: dict[str, Any]) -> list[dict[str, Any]]:
    materials = json.loads(MATERIALS_FILE.read_text(encoding="utf-8"))
    research = load_research(research_dir)
    missing = [m["id"] for m in materials if m["id"] not in research]
    if missing:
        raise SystemExit(f"Research missing for material ids: {missing}")
    unselected = [m["id"] for m in materials if str(m["id"]) not in selection["materials"]]
    if unselected:
        raise SystemExit(f"No equipment selected for material ids: {unselected}")

    namer = FileNamer()
    catalog: list[dict[str, Any]] = []
    for item in materials:
        found = research[item["id"]]
        material_docs = [
            {**doc, "file": namer.assign(doc["url"], material_base(item, doc), MATERIAL_DIR)}
            for doc in found.get("material_docs", [])
            if doc.get("url")
        ]
        equipment_docs = [
            {**doc, "file": namer.assign(doc["url"], equipment_base(doc), EQUIPMENT_DIR)}
            for doc in found.get("equipment_docs", [])
            if doc.get("url")
        ]
        selected = [
            {"key": key, **selection["equipment"][key]}
            for key in selection["materials"][str(item["id"])]
        ]
        catalog.append({**item, "research": found, "material_docs": material_docs,
                        "equipment_docs": equipment_docs, "selected_equipment": selected})
    return catalog


def load_download_status() -> dict[str, str]:
    if not DOWNLOAD_REPORT.exists():
        return {}
    return {row["file"]: row["status"] for row in json.loads(DOWNLOAD_REPORT.read_text())}


def is_shown(doc: dict[str, Any], status: dict[str, str]) -> bool:
    """Main sheets show only real PDFs that downloaded and passed content verification."""
    return bool(doc["file"]) and status.get(doc["file"]) == "ok" and bool(doc.get("verified"))


def shown_material_docs(item: dict[str, Any], status: dict[str, str]) -> list[dict[str, Any]]:
    docs = {d["url"]: d for d in item["material_docs"] if is_shown(d, status)}
    return sorted(docs.values(), key=lambda d: d.get("exact_or_equivalent") != "exact")


def equipment_doc_pool(catalog: list[dict[str, Any]], status: dict[str, str]) -> list[dict[str, Any]]:
    """Every shown equipment PDF in the catalog: a passport belongs to the equipment, not the material."""
    pool = {d["url"]: d for item in catalog for d in item["equipment_docs"] if is_shown(d, status)}
    return list(pool.values())


def shown_equipment_docs(item: dict[str, Any], pool: list[dict[str, Any]]) -> list[tuple[str, dict[str, Any]]]:
    """(equipment names, doc) pairs for the selected equipment, one entry per URL.

    A document covering several selected items (e.g. one brochure for a filling
    machine and its adapter) is listed once, labelled with all their names.
    """
    names: dict[str, list[str]] = {}
    docs: dict[str, dict[str, Any]] = {}
    for equipment in item["selected_equipment"]:
        for doc in pool:
            if any(p in doc["url"] for p in equipment["docs"]):
                names.setdefault(doc["url"], []).append(equipment["name"])
                docs[doc["url"]] = doc
    return [(" / ".join(names[url]), docs[url]) for url in names]


def equipment_names(item: dict[str, Any]) -> str:
    return "\n".join(e["name"] for e in item["selected_equipment"])


def no_doc_text(item: dict[str, Any], kind: str) -> str:
    if kind == "equipment" and all(e["key"] in NO_PASSPORT_KEYS for e in item["selected_equipment"]):
        return "Talab qilinmaydi (qo'lda bajariladi)"
    return "Tekshirilgan PDF topilmadi (Manbalar varag'iga qarang)"


def doc_label(doc: dict[str, Any]) -> str:
    mark = "aniq" if doc.get("exact_or_equivalent") == "exact" else "ekvivalent"
    return f"{doc.get('type', 'DOC')}, {mark}: {doc.get('title', '')}"


def style_sheet(ws, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


HEADERS = ["№", "Material nomi", "Material spetsifikatsiyasi (pasporti)",
           "Material uchun jihoz", "Jihoz pasporti"]
MAIN_WIDTHS = [5, 40, 70, 45, 70]


def material_title(item: dict[str, Any]) -> str:
    return f"{item['name']}\n({item['sap']}; {item['spec']})"


def encode_url(url: str) -> str:
    parts = urlsplit(url)
    return urlunsplit(parts._replace(path=quote(parts.path, safe="/%:@+-._~")))


def write_file_name_sheet(ws, catalog: list[dict[str, Any]], status: dict[str, str]) -> None:
    ws.append(HEADERS)
    pool = equipment_doc_pool(catalog, status)
    for item in catalog:
        mat = shown_material_docs(item, status)
        eq = shown_equipment_docs(item, pool)
        ws.append([
            item["id"],
            material_title(item),
            "\n".join(Path(d["file"]).name for d in mat) or no_doc_text(item, "material"),
            equipment_names(item),
            "\n".join(f"{name}: {Path(d['file']).name}" for name, d in eq) or no_doc_text(item, "equipment"),
        ])
    style_sheet(ws, MAIN_WIDTHS)


def put_link(cell, doc: dict[str, Any], label: str) -> None:
    url = encode_url(doc["url"])
    cell.value = f"{label}\n{url}"
    cell.hyperlink = url
    cell.font = LINK_FONT


def write_link_sheet(ws, catalog: list[dict[str, Any]], status: dict[str, str]) -> None:
    """One material spans several rows: one direct PDF link per cell."""
    ws.append(HEADERS)
    pool = equipment_doc_pool(catalog, status)
    row = 2
    for item in catalog:
        mat = shown_material_docs(item, status)
        eq = shown_equipment_docs(item, pool)
        span = max(len(mat), len(eq), 1)
        ws.cell(row, 1, item["id"])
        ws.cell(row, 2, material_title(item))
        ws.cell(row, 4, equipment_names(item))
        for offset, doc in enumerate(mat):
            put_link(ws.cell(row + offset, 3), doc, doc_label(doc))
        for offset, (name, doc) in enumerate(eq):
            put_link(ws.cell(row + offset, 5), doc, f"{name} — {doc.get('type', '')}: {doc.get('title', '')}")
        if not mat:
            ws.cell(row, 3, no_doc_text(item, "material"))
        if not eq:
            ws.cell(row, 5, no_doc_text(item, "equipment"))
        if span > 1:
            for col in (1, 2, 4):
                ws.merge_cells(start_row=row, start_column=col, end_row=row + span - 1, end_column=col)
        row += span
    style_sheet(ws, MAIN_WIDTHS)
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if cell.hyperlink:
                cell.font = LINK_FONT


def write_sources_sheet(ws, catalog: list[dict[str, Any]], status: dict[str, str]) -> None:
    ws.append(["№", "Material", "Hujjat turi", "Sarlavha (haqiqiy hujjat)", "Fayl", "URL",
               "Aniq/ekvivalent", "Mazmuni tekshirilgan", "Yuklash holati", "Asosiy jadvalda", "Izoh"])
    pool = equipment_doc_pool(catalog, status)
    for item in catalog:
        shown = {d["url"] for d in shown_material_docs(item, status)}
        shown |= {d["url"] for _, d in shown_equipment_docs(item, pool)}
        for kind, docs in (("Material", item["material_docs"]), ("Jihoz", item["equipment_docs"])):
            for doc in docs:
                ws.append([
                    item["id"], item["name"], f"{kind}: {doc.get('type', '')}", doc.get("title", ""),
                    Path(doc["file"]).name if doc["file"] else "(veb-sahifa, PDF emas)", doc["url"],
                    doc.get("exact_or_equivalent", "-"),
                    "ha" if doc.get("verified") else "yo'q",
                    status.get(doc["file"], "yuklanmagan") if doc["file"] else "-",
                    "ha" if doc["url"] in shown else "yo'q",
                    doc.get("note", ""),
                ])
    style_sheet(ws, [5, 30, 16, 60, 45, 60, 14, 12, 14, 12, 60])


def write_identification_sheet(ws, catalog: list[dict[str, Any]]) -> None:
    ws.append(["№", "SAP", "Material", "Spec (manba)", "Qadoq", "Vazifa", "Ishlab chiqaruvchi",
               "Aniqlangan mahsulot", "Ishonch", "Tanlangan jihoz (ishonch)", "Boshqa variantlar", "Izoh"])
    for item in catalog:
        r = item["research"]
        ws.append([
            item["id"], item["sap"], item["name"], item["spec"], item["package"], item["function"],
            r.get("manufacturer", ""), r.get("product_identified", ""), r.get("manufacturer_confidence", ""),
            "\n".join(f"{e['name']} ({e.get('confidence', '?')})" for e in item["selected_equipment"]),
            "\n".join(f"{e.get('vendor', '')} {e.get('model', '')}: {e.get('rationale', '')}"
                      for e in r.get("equipment", [])),
            r.get("notes", ""),
        ])
    style_sheet(ws, [5, 14, 35, 16, 14, 18, 30, 40, 10, 40, 70, 70])


def write_xlsx(catalog: list[dict[str, Any]], target: Path, with_links: bool) -> None:
    status = load_download_status()
    wb = Workbook()
    main = wb.active
    main.title = "Jadval"
    if with_links:
        write_link_sheet(main, catalog, status)
    else:
        write_file_name_sheet(main, catalog, status)
    write_sources_sheet(wb.create_sheet("Manbalar"), catalog, status)
    write_identification_sheet(wb.create_sheet("Identifikatsiya"), catalog)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    catalog = build_catalog(Path(sys.argv[1]), load_selection())
    CATALOG_FILE.write_text(json.dumps(catalog, ensure_ascii=False, indent=1), encoding="utf-8")
    for target, with_links in ((XLSX_FILE, False), (XLSX_LINKS_FILE, True)):
        write_xlsx(catalog, target, with_links)
        print(f"{len(catalog)} materials -> {target.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
