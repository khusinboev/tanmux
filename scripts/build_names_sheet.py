"""Build output/material_equipment_names.xlsx: clean material names + equipment names.

Passport columns are intentionally left empty. Equipment names come from
data/equipment_selection.json (same source as the main workbooks). The "Izoh"
sheet maps every clean name back to the plant list (name, SAP, spec) with
identification and equipment confidence.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output" / "material_equipment_names.xlsx"
SELECTION_FILE = ROOT / "data" / "equipment_selection.json"
CATALOG_FILE = ROOT / "data" / "catalog.json"

# Searchable material names: the confirmed real product name where identified,
# otherwise the plant-list name with obvious typos fixed.
NAMES: dict[int, str] = {
    1: "CHEMFOS 740R L CLEAR",
    2: "ED7010 Binder Gray",
    3: "ED7010 Paste Gray",
    4: "BODY SEALER RSB-430",
    5: "UBC UNDER COAT RSU-230",
    6: "UBS COAT",
    7: "PVC ANTI CHIP",
    8: "CR-100 PRIMER SURFACER LIGHT GRAY (N7.0)",
    9: "CR-100 PRIMER SURFACER PURE WHITE (N9.0)",
    10: "PRIMER TYPE THINNER",
    11: "(GAZ) BASE COAT SUMMIT WHITE (2C1B)",
    12: "(GK2) BASE COAT SON OF A GUN GRAY MET-3",
    13: "(GD8) BASE COAT BURNT COCONUT MET",
    14: "(GB0) BASE COAT BLACK MEET KETTLE MET-3",
    15: "(GYM) BASE COAT SATIN STEEL GRAY MET",
    16: "(GAN) BASE COAT SWITCHBLADE SILVER",
    17: "(GL8) BASE COAT RED-E OR NOT RED MET-3",
    18: "(GUM) BASE COAT SEEKER MET-3",
    19: "(G7G) BASE COAT ZEUS",
    20: "(GZ8) BASE COAT SOME KINDA BLUE",
    21: "(G5J) BASE COAT DARKMOON BLUE",
    22: "(G4V) BASE COAT POW ZINGA MET-3",
    23: "RHT-031 Top Metallic Thinner (BASE TYPE THINNER)",
    24: "TOP CLEAR HITEC 2000",
    25: "CLEAR THINNER",
    26: "RHT-031 Top Metallic Thinner (CLEANING THINNER)",
    27: "CROWN Brake Fluid DOT-4 Plus",
    28: "DEXRON-VI Power Steering Fluid (GMN10060)",
    29: "ANTIFREEZE DX-COOL CONCENTRATE RED (FELIX)",
    30: "Idemitsu MTF LSJ11EP",
    31: "XGP 75W-85 GL-4 Manual Transmission Fluid",
    32: "DEXRON-VI Automatic Transmission Fluid",
    33: "COOLIB R134a (HFC-134a) Refrigerant",
    34: "AI-92 Gasoline",
    35: "Windshield Washer Fluid (Ethanol 80%)",
    36: "Sikaflex-250 UZ-1 / BETASEAL",
    37: "ADHESIVE-HDLNG T/PNL (headliner adhesive)",
    38: "Sika Primer-209 D",
    39: "Sika Primer-206 G+P",
    40: "Sika Aktivator-100",
    41: "RAD 11724 (SIKA-100-UZ)",
    42: "LUBRICANT-WHL PILOT HOLE",
    43: "Klüber MICROLUBE GNY 202",
}

HEADERS = ["Material nomi", "Material spetsifikatsiyasi (pasporti)",
           "Material uchun jihoz", "Jihoz pasporti"]


def style(ws, widths: list[int]) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def main() -> None:
    selection = json.loads(SELECTION_FILE.read_text(encoding="utf-8"))
    catalog = {i["id"]: i for i in json.loads(CATALOG_FILE.read_text(encoding="utf-8"))}
    if sorted(NAMES) != sorted(catalog):
        raise SystemExit("NAMES ids do not match catalog ids")

    wb = Workbook()
    ws = wb.active
    ws.title = "Jadval"
    ws.append(HEADERS)
    notes = wb.create_sheet("Izoh")
    notes.append(["№", "Material nomi", "Zavod ro'yxatidagi nomi", "SAP", "Manbadagi spec",
                  "Identifikatsiya ishonchi", "Jihoz (ishonch)"])
    for item_id in sorted(NAMES):
        item = catalog[item_id]
        keys = selection["materials"][str(item_id)]
        equipment = [selection["equipment"][key] for key in keys]
        ws.append([NAMES[item_id], None, "\n".join(e["name"] for e in equipment), None])
        notes.append([item_id, NAMES[item_id], item["name"], item["sap"], item["spec"],
                      item["research"].get("manufacturer_confidence", ""),
                      "\n".join(f"{e['name']} ({e['confidence']})" for e in equipment)])
    style(ws, [48, 30, 55, 30])
    style(notes, [5, 45, 45, 30, 22, 14, 60])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{len(NAMES)} rows -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
